import os
import uuid
import shutil
import json
from typing import List, Dict, Optional, BinaryIO, Any
from datetime import datetime
import logging
from pathlib import Path
from fastapi import UploadFile, HTTPException, status

from app.models.documents import (
    DocumentCreate,
    DocumentUpdate,
    Document,
    DocumentInDB,
    DocumentType,
    DocumentStatus,
    DocumentMetadata,
    DocumentList
)
from app.core.config import settings

# Import document parsing libraries
import PyPDF2 as pypdf2
import docx2txt
import pandas as pd
import csv
import openpyxl

logger = logging.getLogger(__name__)

# Define the path for storing document data
DOCUMENTS_DIR = Path("backend/data/documents")
DOCUMENTS_DB = DOCUMENTS_DIR / "documents.json"
PROCESSED_DIR = DOCUMENTS_DIR / "processed"


def model_to_dict(model):
    """Convert a Pydantic model to dictionary, working with both v1 and v2"""
    try:
        # For Pydantic v2
        return model.model_dump()
    except AttributeError:
        # For Pydantic v1
        return model.dict()


class DocumentService:
    """Service for document management operations"""
    
    def __init__(self, documents_dir: Optional[Path] = None, documents_db: Optional[Path] = None):
        self.documents_dir = documents_dir or DOCUMENTS_DIR
        self.documents_db = documents_db or DOCUMENTS_DB
        self._ensure_dirs()
    
    def _ensure_dirs(self):
        """Ensure the document directories exist"""
        self.documents_dir.mkdir(parents=True, exist_ok=True)
        
        if not self.documents_db.exists():
            with open(self.documents_db, "w") as f:
                json.dump({"documents": []}, f)
    
    def _read_documents(self) -> List[Dict]:
        """Read documents from the storage file"""
        try:
            with open(self.documents_db, "r") as f:
                data = json.load(f)
                return data.get("documents", [])
        except (FileNotFoundError, json.JSONDecodeError):
            logger.error(f"Error reading documents file {self.documents_db}")
            return []
    
    def _write_documents(self, documents: List[Dict]):
        """Write documents to the storage file"""
        try:
            with open(self.documents_db, "w") as f:
                json.dump({"documents": documents}, f, indent=2)
        except Exception as e:
            logger.error(f"Error writing documents file: {str(e)}")
            raise
    
    def _get_document_path(self, document_id: str) -> Path:
        """Get the document storage path for a specific document ID"""
        return self.documents_dir / document_id
    
    def get_document(self, document_id: str) -> Optional[Document]:
        """Get a document by ID"""
        documents = self._read_documents()
        for doc in documents:
            if doc["id"] == document_id:
                return Document(**doc)
        return None
    
    def get_documents(self, profile_id: Optional[str] = None, skip: int = 0, limit: int = 100) -> DocumentList:
        """Get a list of documents, optionally filtered by profile"""
        documents = self._read_documents()
        
        if profile_id:
            documents = [doc for doc in documents if doc["profile_id"] == profile_id]
        
        total = len(documents)
        documents = documents[skip:skip + limit]
        
        return DocumentList(
            total=total,
            documents=[Document(**doc) for doc in documents]
        )
    
    async def create_document(
        self,
        document_create: DocumentCreate,
        file: UploadFile
    ) -> Document:
        """Create a new document record and save the uploaded file"""
        # Validate file type
        file_ext = os.path.splitext(file.filename or "")[1].lower().strip(".")
        if file_ext not in settings.allowed_extensions_list:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Unsupported file type: {file_ext}. Allowed types: {settings.allowed_extensions_list}"
            )
        
        # Generate document ID and create storage path
        document_id = str(uuid.uuid4())
        document_dir = self._get_document_path(document_id)
        document_dir.mkdir(parents=True, exist_ok=True)
        
        # Save the uploaded file
        file_path = document_dir / f"document.{file_ext}"
        
        try:
            with file_path.open("wb") as f:
                content = await file.read()
                f.write(content)
        except Exception as e:
            logger.error(f"Error saving document file: {str(e)}")
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=f"Failed to save document file: {str(e)}"
            )
        
        # Create document record
        document_data = document_create.dict()
        document = DocumentInDB(
            id=document_id,
            **document_data,
            file_path=str(file_path),
            upload_date=datetime.now(),
            status=DocumentStatus.PENDING
        )
        
        # Save to database
        documents = self._read_documents()
        documents.append(document.dict())
        self._write_documents(documents)
        
        # Process the document asynchronously (in a real app, this would be a background task)
        # For simplicity, we'll do it directly here
        try:
            await self.process_document(document_id)
        except Exception as e:
            logger.error(f"Error processing document: {str(e)}")
            # Mark as failed but still return the document
            self.update_document_status(document_id, DocumentStatus.FAILED, str(e))
        
        return Document(**document.dict())
    
    async def process_document(self, document_id: str) -> Optional[Document]:
        """Process a document to extract text and metadata for AI understanding"""
        document = self.get_document(document_id)
        if not document:
            logger.error(f"Document {document_id} not found for processing")
            return None
            
        # Update status to processing
        self.update_document_status(document_id, DocumentStatus.PROCESSING)
        
        try:
            file_path = Path(document.file_path)
            processed_dir = PROCESSED_DIR / document_id
            processed_dir.mkdir(parents=True, exist_ok=True)
            processed_text_path = processed_dir / "extracted_text.txt"
            processed_metadata_path = processed_dir / "metadata.json"
            
            # Extract text based on document type
            extracted_text = ""
            metadata = DocumentMetadata()
            
            if document.document_type == DocumentType.PDF:
                # Process PDF
                with open(file_path, 'rb') as f:
                    reader = pypdf2.PdfReader(f)
                    metadata.page_count = len(reader.pages)
                    
                    # Extract text from all pages
                    text_parts = []
                    for page_num in range(len(reader.pages)):
                        page = reader.pages[page_num]
                        text_parts.append(page.extract_text())
                    
                    extracted_text = "\n\n".join(text_parts)
                    
                    # Extract additional metadata if available
                    if reader.metadata:
                        metadata.author = reader.metadata.author
                        if reader.metadata.creation_date:
                            metadata.created_date = reader.metadata.creation_date
                        if reader.metadata.modification_date:
                            metadata.modified_date = reader.metadata.modification_date
                    
                    # Count words
                    metadata.word_count = len(extracted_text.split())
                    metadata.size_bytes = file_path.stat().st_size
            
            elif document.document_type == DocumentType.DOCX:
                # Process DOCX
                extracted_text = docx2txt.process(file_path)
                metadata.word_count = len(extracted_text.split())
                metadata.size_bytes = file_path.stat().st_size
            
            elif document.document_type == DocumentType.TXT:
                # Process TXT
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                    extracted_text = f.read()
                metadata.word_count = len(extracted_text.split())
                metadata.size_bytes = file_path.stat().st_size
            
            elif document.document_type == DocumentType.CSV:
                # Process CSV
                df = pd.read_csv(file_path)
                # Convert dataframe to text
                extracted_text = df.to_string()
                metadata.size_bytes = file_path.stat().st_size
                # Add CSV specific metadata
                metadata.page_count = 1
                metadata.word_count = len(extracted_text.split())
            
            elif document.document_type == DocumentType.XLSX:
                # Process XLSX
                df = pd.read_excel(file_path, sheet_name=None)
                texts = []
                for sheet_name, sheet_df in df.items():
                    texts.append(f"Sheet: {sheet_name}\n{sheet_df.to_string()}")
                extracted_text = "\n\n".join(texts)
                metadata.size_bytes = file_path.stat().st_size
                metadata.page_count = len(df)
                metadata.word_count = len(extracted_text.split())
            
            # Save extracted text to file
            with open(processed_text_path, 'w', encoding='utf-8') as f:
                f.write(extracted_text)
            
            # Update document with metadata
            updated_doc = self.update_document_metadata(document_id, metadata)
            if updated_doc:
                # Mark document as completed
                return self.update_document_status(document_id, DocumentStatus.COMPLETED)
            
            return None
            
        except Exception as e:
            logger.error(f"Error processing document {document_id}: {str(e)}")
            # Mark document as failed
            self.update_document_status(document_id, DocumentStatus.FAILED, str(e))
            return None
    
    def update_document(self, document_id: str, document_update: DocumentUpdate) -> Optional[Document]:
        """Update document information"""
        documents = self._read_documents()
        
        for i, doc in enumerate(documents):
            if doc["id"] == document_id:
                # Update document fields
                # For Pydantic v1
                update_data = document_update.dict(exclude_unset=True)
                
                for field, value in update_data.items():
                    if value is not None:
                        doc[field] = value
                
                # Update the documents list
                documents[i] = doc
                self._write_documents(documents)
                
                return Document(**doc)
        
        return None
    
    def update_document_status(
        self,
        document_id: str,
        status: DocumentStatus,
        error: Optional[str] = None
    ) -> Optional[Document]:
        """Update a document's processing status"""
        documents = self._read_documents()
        
        for i, doc in enumerate(documents):
            if doc["id"] == document_id:
                # Update status
                doc["status"] = status
                
                # Set error if provided
                if error:
                    doc["processing_error"] = error
                
                # Update the documents list
                documents[i] = doc
                self._write_documents(documents)
                
                return Document(**doc)
        
        return None
    
    def update_document_metadata(
        self,
        document_id: str,
        metadata: DocumentMetadata
    ) -> Optional[Document]:
        """Update a document's metadata"""
        documents = self._read_documents()
        
        for i, doc in enumerate(documents):
            if doc["id"] == document_id:
                # Update metadata
                doc["metadata"] = metadata.dict()
                
                # Update the documents list
                documents[i] = doc
                self._write_documents(documents)
                
                return Document(**doc)
        
        return None
    
    def delete_document(self, document_id: str) -> bool:
        """Delete a document and its associated files"""
        documents = self._read_documents()
        
        for i, doc in enumerate(documents):
            if doc["id"] == document_id:
                # Remove from the list
                documents.pop(i)
                self._write_documents(documents)
                
                # Delete the document directory
                document_dir = self._get_document_path(document_id)
                if document_dir.exists():
                    shutil.rmtree(document_dir)
                
                return True
        
        return False


# Create a global document service instance
document_service = DocumentService() 